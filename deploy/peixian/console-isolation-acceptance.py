"""Independent live console acceptance using only synthetic test-account data.

Run with the control service's Python environment (httpx/openpyxl/pypdf/python-docx).
No Docker, native Agent access, real model generation, legacy account access,
password changes, or Skill/plugin mutations are performed.

--credentials must name a JSON file under this script's .secrets directory:
{
  "base_url": "http://127.0.0.1:14090",
  "admin": {"username": "admin", "password_file": "console-admin-password"},
  "users": [
    {"username": "console-test", "password_file": "console-test-password"},
    {"username": "console-test-b", "password_file": "console-test-b-password"}
  ]
}
A credential may use "password" instead of "password_file". Password files are
resolved only beneath .secrets. Initial passwords must already have been changed.
Existing Skills and exclusively granted plugins are discovered read-only. Missing
fixtures produce SKIP, never PASS. The report contains no credentials, IDs,
request/response payloads, document text, account names or private URLs.

--self-test generates and validates the synthetic fixtures locally; it never
loads credentials or makes HTTP requests. Live execution cleans up only file,
session and token IDs created by this invocation.
"""
import argparse
import asyncio
from contextlib import AsyncExitStack, suppress
from dataclasses import dataclass, field
from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
import re
import sys
import time
from urllib.parse import urlsplit
import uuid
import zipfile

import httpx
from docx import Document
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject

P = "/api/console/v1"
ROOT = Path(__file__).resolve().parent
LIMIT = 20 * 1024 * 1024
TERMINAL = {"ready", "partial", "no_text", "unsupported", "failed"}


class CheckFailure(Exception):
    """Only developer-authored, non-sensitive error codes may be passed here."""


def need(condition, code):
    if not condition:
        raise CheckFailure(code)


def sha(data):
    return hashlib.sha256(data).hexdigest()


def stamp():
    return datetime.now(timezone.utc).isoformat()


@dataclass
class Credential:
    username: str = field(repr=False)
    password: str = field(repr=False)


@dataclass
class Peer:
    credential: Credential = field(repr=False)
    client: httpx.AsyncClient = field(repr=False)
    user: dict = field(repr=False)
    files: list = field(default_factory=list, repr=False)
    sessions: list = field(default_factory=list, repr=False)
    tokens: list = field(default_factory=list, repr=False)


class Report:
    def __init__(self):
        self.started = stamp()
        self.run = uuid.uuid4().hex
        self.checks = []

    def add(self, name, status, code=""):
        self.checks.append({"name": name, "status": status, "code": code})
        print(f"{status}: {name}" + (f" [{code}]" if code else ""), flush=True)

    async def check(self, name, function):
        try:
            value = await function()
            self.add(name, "PASS")
            return value
        except CheckFailure as error:
            self.add(name, "FAIL", str(error))
        except Exception as error:
            # Never print exception strings: HTTP/library errors may contain paths.
            self.add(name, "FAIL", type(error).__name__)
        return None

    def skip(self, name, code="missing_synthetic_fixture"):
        self.add(name, "SKIP", code)

    def save(self, target):
        target.parent.mkdir(parents=True, exist_ok=True)
        counts = {key: sum(check["status"] == key for check in self.checks) for key in ("PASS", "FAIL", "SKIP")}
        payload = {"run": self.run, "started_at": self.started, "finished_at": stamp(), "counts": counts,
                   "scope": "synthetic_console_accounts_no_model_generation", "checks": self.checks}
        target.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        print("COUNTS: " + json.dumps(counts, sort_keys=True), flush=True)
        return counts


def credentials(path):
    secret_root = (ROOT / ".secrets").resolve()
    path = path.resolve(strict=True)
    if not path.is_relative_to(secret_root):
        raise CheckFailure("credentials_outside_local_secret_directory")
    config = json.loads(path.read_text(encoding="utf-8-sig"))

    def read(item):
        if "password_file" in item:
            password_path = Path(item["password_file"])
            if not password_path.is_absolute():
                password_path = secret_root / password_path
            password_path = password_path.resolve(strict=True)
            need(password_path.is_relative_to(secret_root), "password_file_outside_secret_directory")
            password = password_path.read_text(encoding="utf-8-sig").rstrip("\r\n")
        else:
            password = item["password"]
        need(isinstance(password, str) and bool(password), "invalid_private_credential")
        need(isinstance(item["username"], str), "invalid_private_credential")
        return Credential(item["username"], password)

    base = config.get("base_url", "http://127.0.0.1:14090").rstrip("/")
    parsed = urlsplit(base)
    need(parsed.scheme == "http" and parsed.hostname in ("127.0.0.1", "localhost")
         and not parsed.username and not parsed.password and not parsed.query and not parsed.fragment
         and parsed.path in ("", "/"), "acceptance_requires_loopback_console")
    normal = [read(item) for item in config["users"]]
    need(len(normal) == 2 and normal[0].username != normal[1].username, "need_two_distinct_test_accounts")
    need(all(item.username not in ("client-a", "client-b", "admin") for item in normal), "legacy_accounts_forbidden")
    return base, read(config["admin"]), normal


def new_client(base, **kwargs):
    return httpx.AsyncClient(base_url=base, headers={"Origin": base}, trust_env=False,
                             follow_redirects=False, timeout=httpx.Timeout(90, connect=10), **kwargs)


async def login(client, credential):
    response = await client.post(P + "/auth/login", json={"username": credential.username, "password": credential.password})
    need(response.status_code == 200, "login_failed")
    data = response.json()
    need(isinstance(data.get("csrf_token"), str), "login_contract")
    client.headers["X-CSRF-Token"] = data["csrf_token"]
    user = data["user"]
    need(not user.get("must_change_password"), "initial_password_change_required")
    return user


async def json_request(peer, method, path, expected=200, **kwargs):
    response = await peer.client.request(method, P + path, **kwargs)
    need(response.status_code == expected, "unexpected_http_status")
    return response.json()


def pdf_fixture(marker, *, image_only=False, encrypted=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=200)
    contents = DecodedStreamObject()
    if image_only:
        image = DecodedStreamObject()
        # An actual raster XObject, without a PDF text layer; this is not a blank page.
        image.set_data(bytes(0 if ((x // 4) + (y // 4)) % 2 else 255 for y in range(32) for x in range(64)))
        image.update({NameObject("/Type"): NameObject("/XObject"), NameObject("/Subtype"): NameObject("/Image"),
                      NameObject("/Width"): NumberObject(64), NameObject("/Height"): NumberObject(32),
                      NameObject("/ColorSpace"): NameObject("/DeviceGray"), NameObject("/BitsPerComponent"): NumberObject(8)})
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/XObject"): DictionaryObject({NameObject("/Im0"): writer._add_object(image)})})
        contents.set_data(b"q 200 0 0 100 20 20 cm /Im0 Do Q")
    else:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        contents.set_data(("BT /F1 12 Tf 20 100 Td (" + marker + ") Tj ET").encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(contents)
    if encrypted:
        writer.encrypt("synthetic-document-password")
    result = io.BytesIO()
    writer.write(result)
    return result.getvalue()


def normal_fixtures(marker):
    workbook = Workbook()
    workbook.active.title = "Synthetic"
    workbook.active.append(["item", "value"])
    workbook.active.append([marker, 123])
    xlsx = io.BytesIO()
    workbook.save(xlsx)
    document = Document()
    document.add_paragraph(marker)
    cells = document.add_table(rows=1, cols=2).rows[0].cells
    cells[0].text, cells[1].text = "synthetic", "123"
    docx = io.BytesIO()
    document.save(docx)
    return {
        "txt": ((marker + "\nsecond line").encode(), {"type": "text", "line_start": 1}),
        "md": (("# Synthetic\n" + marker).encode(), {"type": "text", "line_start": 1}),
        "csv": (("item,value\n" + marker + ",123\n").encode(), {"type": "csv", "row_start": 2}),
        "xlsx": (xlsx.getvalue(), {"type": "xlsx", "sheet": "Synthetic", "row_start": 2}),
        "pdf": (pdf_fixture(marker), {"type": "pdf", "page": 1}),
        "docx": (docx.getvalue(), {"type": "docx", "paragraph": 1}),
    }


def expansion_bomb():
    # A valid, highly compressed 201 MiB member. Generation holds only a 1 MiB
    # block; neither this script nor fixture self-test ever decompresses it.
    result = io.BytesIO()
    block = b"0" * (1024 * 1024)
    with zipfile.ZipFile(result, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        with archive.open("padding.xml", "w") as member:
            for _ in range(201):
                member.write(block)
    return result.getvalue()


async def upload(peer, name, content):
    value = await json_request(peer, "POST", "/files", expected=202,
                               files={"file": (name, content, "application/octet-stream")})
    identity = value.get("id")
    need(isinstance(identity, str) and re.fullmatch(r"[0-9a-f]{32}", identity), "opaque_file_id_missing")
    peer.files.append(identity)
    need("source" not in value and "absolute_path" not in value and "path" not in value, "internal_path_exposed")
    return identity


async def parsed(peer, identity, timeout=90):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        value = await json_request(peer, "GET", f"/files/{identity}/text")
        need(set(value) == {"text", "chunks", "truncated", "status", "name"}, "text_contract")
        if value["status"] in TERMINAL:
            return value
        need(value["status"] in ("uploading", "queued", "parsing"), "unknown_parse_status")
        await asyncio.sleep(0.3)
    raise CheckFailure("parse_timeout")


async def file_round_trip(peer, report, extension, content, marker, expected_source):
    identity = await upload(peer, f"acceptance-{report.run[:12]}.{extension}", content)
    value = await parsed(peer, identity)
    need(value["status"] == "ready" and marker in value["text"], "text_extraction_failed")
    need(any(marker in chunk["text"] and all(chunk.get("source", {}).get(key) == expected
         for key, expected in expected_source.items()) for chunk in value["chunks"]), "source_provenance_missing")
    if extension == "docx":
        need(any(chunk.get("source", {}).get("table") == 1 for chunk in value["chunks"]), "docx_table_source_missing")
    response = await peer.client.get(P + f"/files/{identity}/download")
    need(response.status_code == 200 and sha(response.content) == sha(content), "download_bytes_changed")
    need(response.headers.get("x-content-type-options") == "nosniff", "download_nosniff_missing")
    preview = await json_request(peer, "GET", f"/files/{identity}/preview")
    need(marker in preview["text"], "preview_missing")
    return identity


async def invalid_file(peer, name, content, expected_status, expected_error=None):
    identity = await upload(peer, name, content)
    value = await parsed(peer, identity)
    need(value["status"] == expected_status, "wrong_invalid_file_status")
    if expected_status in ("failed", "no_text"):
        need(value["text"] == "" and value["chunks"] == [], "invalid_file_exposed_text")
    if expected_error:
        listing = await json_request(peer, "GET", "/files")
        item = next((item for item in listing["items"] if item["id"] == identity), None)
        need(item is not None and item.get("error") == expected_error, "wrong_parse_error_category")
    response = await peer.client.get(P + f"/files/{identity}/download")
    need(response.status_code == 200 and sha(response.content) == sha(content), "invalid_original_not_preserved")
    return identity


async def expect_status(client, method, path, expected, **kwargs):
    response = await client.request(method, P + path, **kwargs)
    need(response.status_code == expected, "unexpected_http_status")


async def create_session(peer, title):
    value = await json_request(peer, "POST", "/sessions", json={"title": title})
    identity = value.get("id")
    need(isinstance(identity, str), "session_contract")
    peer.sessions.append(identity)
    need(set(value) <= {"id", "title", "time", "parentID"}, "internal_session_fields_exposed")
    await json_request(peer, "GET", f"/sessions/{identity}/messages")
    return identity


async def stream_revocation(base, credential, *, bearer=False):
    async with new_client(base) as owner, new_client(base) as stream_client, new_client(base) as replay:
        await login(owner, credential)
        token_id = None
        if bearer:
            token = await owner.post(P + "/tokens", json={"name": "synthetic-isolation-sse"})
            need(token.status_code == 200, "token_creation_failed")
            value = token.json()
            token_id = value["item"]["id"]
            stream_client.headers["Authorization"] = "Bearer " + value["token"]
            replay.headers["Authorization"] = stream_client.headers["Authorization"]
        else:
            # Use the same session cookie on a separate connection. Logging it out
            # must close the already-open SSE response, not only future requests.
            stream_client.cookies.update(owner.cookies)
            replay.cookies.update(owner.cookies)
        need((await replay.get(P + "/me")).status_code == 200, "pre_revocation_auth_failed")
        alive = asyncio.Event()
        revoked = False
        ended_before_revocation = False

        async def observe():
            nonlocal ended_before_revocation
            try:
                async with stream_client.stream("GET", P + "/events", timeout=None) as response:
                    need(response.status_code == 200, "sse_connection_failed")
                    async for line in response.aiter_lines():
                        if line.startswith(": heartbeat"):
                            alive.set()
            finally:
                ended_before_revocation = not revoked

        task = asyncio.create_task(observe())
        try:
            await asyncio.wait_for(alive.wait(), timeout=8)
            need(not task.done(), "sse_not_live_before_revocation")
            revoked = True
            if bearer:
                response = await owner.delete(P + "/tokens/" + token_id)
            else:
                response = await owner.post(P + "/auth/logout")
            need(response.status_code == 200, "revocation_failed")
            need((await replay.get(P + "/me")).status_code == 401, "old_auth_still_accepted")
            await asyncio.wait_for(asyncio.shield(task), timeout=5)
            need(not ended_before_revocation, "sse_ended_before_revocation")
        finally:
            task.cancel()
            with suppress(asyncio.CancelledError, Exception):
                await task
            if bearer and token_id:
                with suppress(Exception):
                    await owner.delete(P + "/tokens/" + token_id)
                with suppress(Exception):
                    await owner.post(P + "/auth/logout")


async def cleanup(peer, report, label):
    for identity in reversed(peer.files):
        async def remove_file(identity=identity):
            deadline = time.monotonic() + 65
            while True:
                response = await peer.client.delete(P + "/files/" + identity)
                if response.status_code in (200, 404):
                    return
                if response.status_code != 409 or time.monotonic() >= deadline:
                    raise CheckFailure("synthetic_file_cleanup_failed")
                await asyncio.sleep(0.5)
        await report.check(label + ".cleanup.file", remove_file)
    for identity in reversed(peer.sessions):
        await report.check(label + ".cleanup.session", lambda identity=identity:
                           expect_status(peer.client, "DELETE", "/sessions/" + identity, 200))
    for identity in peer.tokens:
        await report.check(label + ".cleanup.token", lambda identity=identity:
                           expect_status(peer.client, "DELETE", "/tokens/" + identity, 200))


async def live(config_path, report):
    base, admin_credential, normal_credentials = credentials(config_path)
    async with AsyncExitStack() as stack:
        admin_client = await stack.enter_async_context(new_client(base))
        anonymous = await stack.enter_async_context(new_client(base))
        admin_user = await login(admin_client, admin_credential)
        need(admin_user["role"] == "admin", "administrator_role_missing")
        peers = []
        for credential in normal_credentials:
            client = await stack.enter_async_context(new_client(base))
            user = await login(client, credential)
            need(user["role"] == "user", "test_account_role_invalid")
            runtime = user.get("runtime") or {}
            need(runtime.get("status") == "ready" and runtime.get("revision") == runtime.get("desired"),
                 "test_runtime_not_ready")
            peers.append(Peer(credential, client, user))
        need(peers[0].user["id"] != peers[1].user["id"], "accounts_not_distinct")
        report.add("preflight.two_independent_ready_test_accounts", "PASS")
        samples = [{}, {}]
        try:
            for path in ("/me", "/files", "/sessions", "/skills", "/plugins", "/events"):
                await report.check("auth.anonymous." + path[1:],
                                   lambda path=path: expect_status(anonymous, "GET", path, 401))
            for index, peer in enumerate(peers):
                label = f"account_{index + 1}"
                await report.check(label + ".admin_api_forbidden", lambda peer=peer:
                                   expect_status(peer.client, "GET", "/admin/users", 403))
                async def no_csrf(peer=peer):
                    headers = {key: value for key, value in peer.client.headers.items() if key.lower() != "x-csrf-token"}
                    request = peer.client.build_request("POST", P + "/sessions", json={"title": "not-created"})
                    request.headers.pop("x-csrf-token", None)
                    response = await peer.client.send(request)
                    need(response.status_code == 403, "csrf_missing_was_accepted")
                await report.check(label + ".missing_csrf_forbidden", no_csrf)
                marker = f"PX-{report.run[:12]}-{index + 1}"
                for extension, (content, source) in normal_fixtures(marker).items():
                    identity = await report.check(label + ".file." + extension + ".parse_source_download",
                        lambda extension=extension, content=content, source=source, peer=peer, marker=marker:
                        file_round_trip(peer, report, extension, content, marker, source))
                    if identity:
                        samples[index][extension] = identity
                session = await report.check(label + ".session.create_empty",
                    lambda peer=peer, label=label: create_session(peer, "acceptance-" + report.run[:12] + "-" + label))
                if session:
                    samples[index]["session"] = session
                    # Empty text is a second fail-closed guard: even if invalid
                    # model selection regresses, this cannot be a valid prompt.
                    await report.check(label + ".unapproved_model_forbidden", lambda peer=peer, session=session:
                        expect_status(peer.client, "POST", f"/sessions/{session}/messages", 403,
                                      json={"model_id": "never-authorized-" + report.run, "text": ""}))
            a = peers[0]
            async def duplicates():
                name = "acceptance-" + report.run[:12] + "-same.txt"
                identities = [await upload(a, name, value) for value in (b"first synthetic value", b"second synthetic value")]
                need(identities[0] != identities[1], "duplicate_name_reused_id")
                for identity, value in zip(identities, (b"first synthetic value", b"second synthetic value")):
                    await parsed(a, identity)
                    response = await a.client.get(P + f"/files/{identity}/download")
                    need(response.status_code == 200 and response.content == value, "duplicate_overwrote_source")
            await report.check("files.duplicate_names_preserve_both", duplicates)
            invalid = (
                ("image_pdf", ".pdf", pdf_fixture("synthetic", image_only=True), "no_text", None),
                ("encrypted_pdf", ".pdf", pdf_fixture("synthetic", encrypted=True), "failed", "encrypted_document"),
                ("corrupt_pdf", ".pdf", b"synthetic invalid PDF", "failed", "parse_error"),
                ("corrupt_xlsx", ".xlsx", b"synthetic invalid workbook", "failed", "parse_error"),
                ("corrupt_docx", ".docx", b"synthetic invalid document", "failed", "parse_error"),
                ("zip_expansion_bomb", ".xlsx", expansion_bomb(), "failed", "resource_limit"),
                ("unsupported", ".bin", b"synthetic unsupported data", "unsupported", None),
            )
            for name, extension, content, status, error in invalid:
                await report.check("files." + name, lambda name=name, extension=extension, content=content, status=status, error=error:
                    invalid_file(a, "acceptance-" + report.run[:12] + "-" + name + extension, content, status, error))
            await report.check("files.over_20_mib_rejected", lambda:
                expect_status(a.client, "POST", "/files", 413,
                              files={"file": ("acceptance-oversize.txt", b"x" * (LIMIT + 1), "text/plain")}))
            for filename in ("../outside.txt", "/outside.txt", "..\\outside.txt"):
                await report.check("files.path_filename_rejected", lambda filename=filename:
                    expect_status(a.client, "POST", "/files", 400,
                                  files={"file": (filename, b"synthetic", "text/plain")}))
            for index, peer in enumerate(peers):
                other = 1 - index
                label = f"account_{index + 1}.foreign"
                foreign_file, foreign_session = samples[other].get("txt"), samples[other].get("session")
                if foreign_file:
                    for action in ("text", "preview", "download"):
                        await report.check(label + ".file." + action, lambda action=action, peer=peer, fid=foreign_file:
                            expect_status(peer.client, "GET", f"/files/{fid}/{action}", 404))
                    await report.check(label + ".file.delete", lambda peer=peer, fid=foreign_file:
                        expect_status(peer.client, "DELETE", f"/files/{fid}", 404))
                else:
                    report.skip(label + ".file_ids")
                if foreign_session:
                    await report.check(label + ".session.messages", lambda peer=peer, sid=foreign_session:
                        expect_status(peer.client, "GET", f"/sessions/{sid}/messages", 404))
                    await report.check(label + ".session.update", lambda peer=peer, sid=foreign_session:
                        expect_status(peer.client, "PATCH", f"/sessions/{sid}", 404, json={"unexpected": True}))
                    async def routing(peer=peer, sid=foreign_session, other=other):
                        value = await json_request(peer, "GET", "/sessions", params={"uid": peers[other].user["id"], "directory": "/other"},
                                                   headers={"X-Peixian-Account": peers[other].user["id"]})
                        need(sid not in {item["id"] for item in value["items"]}, "caller_selected_foreign_account")
                    await report.check(label + ".caller_cannot_select_account", routing)
                else:
                    report.skip(label + ".session_ids")
            skills = [(await json_request(peer, "GET", "/skills"))["items"] for peer in peers]
            plugins = [(await json_request(peer, "GET", "/plugins"))["items"] for peer in peers]
            exclusive_plugin_cases = 0
            for index, peer in enumerate(peers):
                other = 1 - index
                label = f"account_{index + 1}.foreign"
                if skills[other]:
                    sid = skills[other][0]["id"]
                    need(sid not in {item["id"] for item in skills[index]}, "foreign_skill_listed")
                    await report.check(label + ".skill.test", lambda peer=peer, sid=sid:
                        expect_status(peer.client, "POST", f"/skills/{sid}/test", 404, json={}))
                    await report.check(label + ".skill.update", lambda peer=peer, sid=sid:
                        expect_status(peer.client, "PATCH", f"/skills/{sid}", 404, json={"unexpected": True}))
                else:
                    report.skip(label + ".skill_ids")
                own_plugins = {item["id"] for item in plugins[index]}
                exclusive = next((item for item in plugins[other] if item["id"] not in own_plugins), None)
                if exclusive:
                    exclusive_plugin_cases += 1
                    await report.check(label + ".plugin.install", lambda peer=peer, item=exclusive:
                        expect_status(peer.client, "PUT", "/plugins/" + item["id"], 404,
                                      json={"version": item["version"], "config": []}))
                    await report.check(label + ".plugin.test", lambda peer=peer, item=exclusive:
                        expect_status(peer.client, "POST", "/plugins/" + item["id"] + "/test", 409, json={}))
            if not exclusive_plugin_cases:
                report.skip("plugins.exclusive_plugin_fixture")
            for path in ("/files", "/sessions", "/skills", "/plugins", "/models", "/results", "/events"):
                await report.check("admin.business_forbidden." + path[1:],
                                   lambda path=path: expect_status(admin_client, "GET", path, 403))
            if samples[0].get("txt"):
                await report.check("admin.account_file_download_forbidden", lambda:
                    expect_status(admin_client, "GET", "/files/" + samples[0]["txt"] + "/download", 403))
            if samples[0].get("session"):
                await report.check("admin.account_messages_forbidden", lambda:
                    expect_status(admin_client, "GET", "/sessions/" + samples[0]["session"] + "/messages", 403))
            await report.check("auth.logout_closes_existing_sse_and_rejects_cookie",
                               lambda: stream_revocation(base, normal_credentials[0], bearer=False))
            await report.check("auth.token_revoke_closes_existing_sse_and_rejects_token",
                               lambda: stream_revocation(base, normal_credentials[1], bearer=True))
        finally:
            for index, peer in enumerate(peers):
                await cleanup(peer, report, f"account_{index + 1}")
                await report.check(f"account_{index + 1}.cleanup.login",
                                   lambda peer=peer: expect_status(peer.client, "POST", "/auth/logout", 200))
            await report.check("admin.cleanup.login", lambda:
                               expect_status(admin_client, "POST", "/auth/logout", 200))


def self_test(report):
    fixtures = normal_fixtures("PX-SYNTHETIC")
    need(len(fixtures) == 6, "fixture_count")
    need("PX-SYNTHETIC" in PdfReader(io.BytesIO(fixtures["pdf"][0])).pages[0].extract_text(), "pdf_fixture")
    need(load_workbook(io.BytesIO(fixtures["xlsx"][0]), read_only=True).active.cell(2, 1).value == "PX-SYNTHETIC", "xlsx_fixture")
    need(Document(io.BytesIO(fixtures["docx"][0])).paragraphs[0].text == "PX-SYNTHETIC", "docx_fixture")
    image = PdfReader(io.BytesIO(pdf_fixture("unused", image_only=True))).pages[0]
    need(not image.extract_text() and bool(image["/Resources"]["/XObject"]), "image_pdf_fixture")
    need(PdfReader(io.BytesIO(pdf_fixture("synthetic", encrypted=True))).is_encrypted, "encrypted_fixture")
    bomb = expansion_bomb()
    with zipfile.ZipFile(io.BytesIO(bomb)) as archive:
        need(sum(item.file_size for item in archive.infolist()) > 200 * 1024 * 1024 and len(bomb) < LIMIT, "bomb_fixture")
    report.add("offline.synthetic_fixture_generation_and_format_validation", "PASS")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--credentials", type=Path)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--self-test", action="store_true")
    args = parser.parse_args()
    report = Report()
    try:
        if args.self_test:
            self_test(report)
        else:
            need(args.credentials is not None, "credentials_argument_required")
            asyncio.run(live(args.credentials, report))
    except CheckFailure as error:
        report.add("run.precondition_or_fixture", "FAIL", str(error))
    except Exception as error:
        report.add("run.unexpected", "FAIL", type(error).__name__)
    target = args.report or ROOT / "output" / ("console-isolation-" + report.run + ".json")
    counts = report.save(target)
    return 1 if counts["FAIL"] else 2 if counts["SKIP"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
