"""Document parser subprocess. It emits bounded JSON and never performs OCR."""

from contextlib import redirect_stderr, redirect_stdout
import csv
from datetime import date, datetime
import io
import json
import os
from pathlib import Path
import sys
import zipfile


MAX_INPUT = 20 * 1024 * 1024
MAX_EXPANDED = 200 * 1024 * 1024
MAX_TEXT = 1_000_000
MAX_CHUNKS = 10000


class ParseLimit(ValueError):
    pass


def archive_check(path):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(entry.file_size for entry in entries) > MAX_EXPANDED:
            raise ParseLimit("archive_limit")
        for entry in entries:
            name = entry.filename.replace("\\", "/")
            if entry.flag_bits & 1 or name.startswith("/") or ".." in name.split("/") or ":" in name:
                raise ParseLimit("unsafe_archive")


class Collector:
    def __init__(self):
        self.chunks = []
        self.length = 0
        self.truncated = False

    def add(self, text, source):
        text = str(text).strip()
        if not text:
            return True
        remaining = MAX_TEXT - self.length
        if remaining <= 0 or len(self.chunks) >= MAX_CHUNKS:
            self.truncated = True
            return False
        if len(text) > remaining:
            text = text[:remaining]
            self.truncated = True
        self.chunks.append({"text": text, "source": source})
        self.length += len(text) + 2
        return not self.truncated

    def result(self, warnings=None):
        text = "\n\n".join(chunk["text"] for chunk in self.chunks)
        return {
            "text": text, "chunks": self.chunks, "truncated": self.truncated,
            "status": "partial" if self.truncated else ("ready" if text else "no_text"),
            "warnings": warnings or [],
        }


def decode_text(data):
    encodings = ("utf-8-sig", "gb18030")
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings = ("utf-16",)
    for encoding in encodings:
        try:
            text = data.decode(encoding)
            if "\0" in text:
                raise ValueError("binary_text")
            return text
        except UnicodeDecodeError:
            continue
    raise ValueError("text_encoding")


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def parse(path, suffix):
    if path.stat().st_size > MAX_INPUT:
        raise ParseLimit("input_limit")
    result = Collector()
    if suffix in (".txt", ".md"):
        lines = decode_text(path.read_bytes()).splitlines()
        for start in range(0, len(lines), 100):
            if not result.add("\n".join(lines[start:start + 100]), {
                "type": "text", "line_start": start + 1, "line_end": min(start + 100, len(lines)),
            }):
                break
        return result.result()
    if suffix == ".csv":
        text = decode_text(path.read_bytes())
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        csv.field_size_limit(MAX_INPUT)
        for row_number, row in enumerate(csv.reader(io.StringIO(text), dialect), 1):
            if not any(value.strip() for value in row):
                continue
            if not result.add(" | ".join(row), {"type": "csv", "row_start": row_number, "row_end": row_number}):
                break
        return result.result()
    if suffix == ".xlsx":
        archive_check(path)
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in workbook.worksheets:
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    if not any(value is not None and str(value).strip() for value in row):
                        continue
                    if not result.add(" | ".join(cell_text(value) for value in row), {
                        "type": "xlsx", "sheet": sheet.title, "row_start": row_number, "row_end": row_number,
                    }):
                        return result.result(["cached_formula_values_only"])
        finally:
            workbook.close()
        return result.result(["cached_formula_values_only"])
    if suffix == ".pdf":
        from pypdf import PdfReader
        document = PdfReader(path, strict=False)
        if document.is_encrypted:
            raise ValueError("encrypted_pdf")
        for page_number, page in enumerate(document.pages, 1):
            if not result.add(page.extract_text() or "", {"type": "pdf", "page": page_number}):
                break
        return result.result(["text_only_no_ocr"])
    if suffix == ".docx":
        archive_check(path)
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document = Document(path)
        paragraph_number = 0
        table_number = 0
        for element in document.element.body.iterchildren():
            if element.tag == qn("w:p"):
                paragraph_number += 1
                paragraph = Paragraph(element, document)
                if not result.add(paragraph.text, {"type": "docx", "paragraph": paragraph_number}):
                    break
            if element.tag == qn("w:tbl"):
                table_number += 1
                table = Table(element, document)
                for row_number, row in enumerate(table.rows, 1):
                    if not result.add(" | ".join(cell.text for cell in row.cells), {
                        "type": "docx", "table": table_number, "row_start": row_number, "row_end": row_number,
                    }):
                        return result.result(["body_text_and_tables"])
        return result.result(["body_text_and_tables"])
    return {"text": "", "chunks": [], "truncated": False, "status": "unsupported", "warnings": []}


def apply_limits():
    if os.name != "posix":
        return
    import resource
    resource.setrlimit(resource.RLIMIT_AS, (384 * 1024 * 1024, 384 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_CPU, (60, 60))
    resource.setrlimit(resource.RLIMIT_FSIZE, (200 * 1024 * 1024, 200 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_NOFILE, (64, 64))
    resource.setrlimit(resource.RLIMIT_CORE, (0, 0))


def main():
    apply_limits()
    try:
        if len(sys.argv) != 3:
            raise ValueError("invalid_arguments")
        with open(os.devnull, "w") as sink, redirect_stdout(sink), redirect_stderr(sink):
            result = parse(Path(sys.argv[1]), sys.argv[2])
    except (ParseLimit, MemoryError):
        result = {"text": "", "chunks": [], "truncated": False, "status": "failed", "error": "resource_limit"}
    except Exception as error:
        code = {"encrypted_pdf": "encrypted_document", "text_encoding": "text_encoding", "binary_text": "text_encoding"}.get(str(error), "parse_error")
        result = {"text": "", "chunks": [], "truncated": False, "status": "failed", "error": code}
    sys.stdout.write(json.dumps(result, ensure_ascii=False, allow_nan=False))


if __name__ == "__main__":
    main()