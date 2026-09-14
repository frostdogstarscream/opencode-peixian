import json
import os
from pathlib import Path
import subprocess
import sys
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from gateway import parser


def test_text_csv_sources_and_encodings(tmp_path):
    path = tmp_path / "chinese.txt"
    path.write_bytes("案件测试\n第二行".encode("gb18030"))
    result = parser.parse(path, ".txt")
    assert result["status"] == "ready"
    assert "案件测试" in result["text"]
    assert result["chunks"][0]["source"]["line_end"] == 2
    path = tmp_path / "data.csv"
    path.write_text('name,value\n"a,b",10\n,,\n', encoding="utf-8")
    result = parser.parse(path, ".csv")
    assert result["chunks"][1]["text"] == "a,b | 10"
    assert result["chunks"][1]["source"]["row_start"] == 2


def test_xlsx_cached_values_and_sheet_rows(tmp_path):
    path = tmp_path / "data.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Evidence"
    sheet.append(["name", "amount"])
    sheet.append(["synthetic", 12])
    sheet.append([None, None])
    sheet.cell(4, 1, "end")
    book.save(path)
    result = parser.parse(path, ".xlsx")
    assert result["status"] == "ready"
    assert result["chunks"][1]["source"] == {"type": "xlsx", "sheet": "Evidence", "row_start": 2, "row_end": 2}
    assert result["chunks"][1]["text"] == "synthetic | 12"
    assert len(result["chunks"]) == 3
    assert "cached_formula_values_only" in result["warnings"]


def test_docx_paragraph_and_table_sources(tmp_path):
    path = tmp_path / "data.docx"
    document = Document()
    document.add_paragraph("Synthetic paragraph")
    cells = document.add_table(rows=1, cols=2).rows[0].cells
    cells[0].text, cells[1].text = "item", "42"
    document.save(path)
    result = parser.parse(path, ".docx")
    assert result["chunks"][0]["source"] == {"type": "docx", "paragraph": 1}
    assert result["chunks"][1]["source"]["table"] == 1
    assert result["chunks"][1]["text"] == "item | 42"


def test_pdf_text_page_and_image_only_status(tmp_path):
    path = tmp_path / "text.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=200)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 20 100 Td (Synthetic PDF text) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    writer.write(path)
    result = parser.parse(path, ".pdf")
    assert "Synthetic PDF text" in result["text"]
    assert result["chunks"][0]["source"] == {"type": "pdf", "page": 1}
    blank = PdfWriter()
    blank.add_blank_page(width=100, height=100)
    blank.write(path)
    result = parser.parse(path, ".pdf")
    assert result["status"] == "no_text"
    assert result["warnings"] == ["text_only_no_ocr"]


def test_zip_traversal_expansion_and_output_limits(tmp_path, monkeypatch):
    path = tmp_path / "bad.xlsx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("../escape.xml", "payload")
    with pytest.raises(parser.ParseLimit):
        parser.parse(path, ".xlsx")
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("small.xml", "payload")
    monkeypatch.setattr(parser, "MAX_EXPANDED", 3)
    with pytest.raises(parser.ParseLimit):
        parser.archive_check(path)
    monkeypatch.setattr(parser, "MAX_TEXT", 10)
    collector = parser.Collector()
    assert not collector.add("a" * 20, {"type": "text"})
    assert collector.result()["status"] == "partial"
    assert len(collector.result()["text"]) == 10


def test_parser_subprocess_emits_only_protocol_and_generic_error(tmp_path):
    path = tmp_path / "bad.pdf"
    path.write_bytes(b"not a PDF, no secrets")
    completed = subprocess.run([sys.executable, "-I", "-B", str(Path(parser.__file__)), str(path), ".pdf"],
                               capture_output=True, timeout=15)
    assert completed.returncode == 0
    result = json.loads(completed.stdout)
    assert result["status"] == "failed" and result["error"] == "parse_error"
    assert completed.stderr == b""
    assert str(path).encode() not in completed.stdout
