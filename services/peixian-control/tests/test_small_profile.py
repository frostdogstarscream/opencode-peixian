import asyncio
import builtins
import importlib.util
import json
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import httpx
import pytest

from examples.console_client import ConsoleClient

spec = importlib.util.spec_from_file_location("small_profile_contract", Path(__file__).parents[1] / "benchmarks/small_profile.py")
profile = importlib.util.module_from_spec(spec)
spec.loader.exec_module(profile)


def manifest():
    return {"base_url": "https://synthetic.invalid", "deployment_id": "loadtest-profile",
            "users": [{"username": f"loadtest-{index}", "token": f"px_synthetic-{index}"} for index in range(2)]}


@pytest.mark.parametrize("invalid", ["normal-user", "duplicate-user", "duplicate-token", "too-many", "no-users",
                                     "non-string", "non-synthetic", "forward-url", "public-http", "token-newline"])
def test_manifest_rejects_unsafe_scope_and_credentials(tmp_path, invalid):
    value = manifest()
    if invalid == "normal-user": value["users"][0]["username"] = "client-a"
    if invalid == "duplicate-user": value["users"][1]["username"] = value["users"][0]["username"]
    if invalid == "duplicate-token": value["users"][1]["token"] = value["users"][0]["token"]
    if invalid == "too-many": value["users"].append({"username": "loadtest-3", "token": "px_synthetic-3"})
    if invalid == "no-users": value["users"] = []
    if invalid == "non-string": value["users"][0]["username"] = None
    if invalid == "non-synthetic": value["deployment_id"] = "production"
    if invalid == "forward-url": value["base_url"] += "/proxy?url=https://unapproved.invalid"
    if invalid == "public-http": value["base_url"] = "http://synthetic.invalid"
    if invalid == "token-newline": value["users"][0]["token"] += "\nInjected: true"
    source = tmp_path / "credentials.json"
    source.write_text(json.dumps(value))
    with pytest.raises(ValueError):
        profile.load_manifest(source)


def args(tmp_path, scenario="plugin-test"):
    source = tmp_path / "credentials.json"
    source.write_text(json.dumps(manifest()))
    return SimpleNamespace(credentials=source, accounts=2, scenario=scenario, model_id="test-model", ca_file=None,
                           synthetic_deployment=True, task_timeout=10, observe_seconds=1, file_type="txt", plugin_version="1.0.0")


@pytest.mark.parametrize("failure", ["overload", "unknown", "network"])
def test_failed_plugin_submission_is_once_per_user_and_report_has_no_secrets(tmp_path, monkeypatch, failure):
    requests = []
    async def respond(request):
        index = request.headers["Authorization"][-1]
        if request.url.path.endswith("/me"):
            return httpx.Response(200, json={"user": {"username": "loadtest-" + index, "role": "user",
                                                      "must_change_password": False, "runtime": {"status": "ready"}}})
        if request.url.path.endswith("/plugins"):
            return httpx.Response(200, json={"items": [{"id": "sample-records", "installed": {
                "enabled": True, "version": "1.0.0", "state": "active"}}]})
        requests.append(request)
        if failure == "network":
            raise httpx.ReadTimeout("https://private.invalid px_leaked-body", request=request)
        return httpx.Response(503 if failure == "overload" else 504,
                              headers={"Retry-After": "1"}, json={"message": "private body"})
    monkeypatch.setattr(profile, "ConsoleClient", lambda *a, **kw: ConsoleClient(*a, **kw, transport=httpx.MockTransport(respond)))
    result = asyncio.run(profile.execute(args(tmp_path)))
    assert len(requests) == 2
    assert all(request.method == "POST" for request in requests)
    assert result["status"] == "failed"
    assert all(item["error"]["result_unknown"] == (failure != "overload") for item in result["stages"][0]["accounts"])
    encoded = json.dumps(result)
    assert not any(value in encoded for value in ("https://", "private body", "px_", "loadtest-0", "loadtest-1"))


def test_all_accounts_are_checked_before_any_mutation(tmp_path, monkeypatch):
    requests = []
    async def respond(request):
        requests.append(request)
        if request.url.path.endswith("/me"):
            return httpx.Response(200, json={"user": {"username": "wrong-user", "role": "user", "runtime": {"status": "ready"}}})
        raise AssertionError("No mutation allowed")
    monkeypatch.setattr(profile, "ConsoleClient", lambda *a, **kw: ConsoleClient(*a, **kw, transport=httpx.MockTransport(respond)))
    with pytest.raises(ValueError):
        asyncio.run(profile.execute(args(tmp_path)))
    assert all(request.method == "GET" for request in requests)


def test_opt_in_required_even_for_programmatic_execution(tmp_path):
    options = args(tmp_path)
    options.synthetic_deployment = False
    with pytest.raises(ValueError, match="synthetic_opt_in_required"):
        asyncio.run(profile.execute(options))


def test_tool_evidence_must_be_new_completed_and_match_fixed_sample_outputs():
    old = {"info": {"id": "old", "role": "assistant", "time": {"completed": 1}}, "parts": [
        {"type": "tool", "tool": "调用已启用的插件", "state": {"status": "completed"},
         "details": {"outputs": {"source": "示例资料服务", "version": "1.0.0", "count": 1}}}]}
    values = [old, {"info": {"id": "new", "role": "assistant", "time": {"completed": 2}}, "parts": []}]
    assert profile.new_message_evidence(values, {"old"}, "1.0.0")["matching_sample_records_outputs"] == 0
    new = json.loads(json.dumps(old))
    new["info"]["id"] = "new"
    assert profile.new_message_evidence([old, new], {"old"}, "1.0.0")["matching_sample_records_outputs"] == 1
    new["parts"][0]["state"]["status"] = "error"
    assert profile.new_message_evidence([old, new], {"old"}, "1.0.0")["matching_sample_records_outputs"] == 0
    new["parts"][0]["state"]["status"] = "completed"
    assert profile.new_message_evidence([old, new], {"old"}, "2.0.0")["matching_sample_records_outputs"] == 0


@pytest.mark.parametrize("failure", ["overload", "unknown", "network"])
def test_model_submission_is_not_replayed_by_workflow(tmp_path, failure):
    calls = []
    class FailingClient:
        async def create_session(self, title): return {"id": "synthetic-session"}
        async def messages(self, session_id): return []
        async def run_message(self, *a, **kw):
            calls.append(1)
            raise profile.ConsoleError("private body", status=503 if failure == "overload" else 504,
                                       result_unknown=failure != "overload")
    result = asyncio.run(profile.workflow(FailingClient(), "answers", args(tmp_path), 1))
    assert calls == [1]
    assert result["status"] == "failed"
    assert "private body" not in json.dumps(result)


def test_file_upload_uses_only_generated_synthetic_text_and_is_not_replayed(tmp_path):
    files = []
    class FailingClient:
        async def upload(self, path):
            files.append(path.read_text(encoding="utf-8"))
            raise profile.ConsoleError("unknown", status=504, result_unknown=True)
    result = asyncio.run(profile.workflow(FailingClient(), "files", args(tmp_path), 1))
    assert len(files) == 1
    assert files[0].startswith("LOADTEST-SYNTHETIC-")
    assert "合成行 099" in files[0]
    assert result["status"] == "failed"
    assert "LOADTEST-SYNTHETIC-" not in json.dumps(result)


def test_minimal_docx_needs_no_native_dependencies_and_real_parser_preserves_text(tmp_path, monkeypatch):
    path = tmp_path / "synthetic.docx"
    lines = ["LOADTEST-SYNTHETIC-marker", "  合成 <段落> & 特殊字符  ", "合成行 099"]
    original_import = builtins.__import__
    def without_native_parser(name, *args, **kwargs):
        if name.split(".")[0] in ("docx", "lxml"):
            raise ModuleNotFoundError("Load client has no document parser")
        return original_import(name, *args, **kwargs)
    with monkeypatch.context() as context:
        context.setattr(builtins, "__import__", without_native_parser)
        profile.write_synthetic_docx(path, lines)
    with ZipFile(path) as archive:
        assert set(archive.namelist()) == {"[Content_Types].xml", "_rels/.rels", "word/document.xml"}
        assert archive.testzip() is None
    from docx import Document
    assert [paragraph.text for paragraph in Document(path).paragraphs] == lines


def test_docx_workflow_uploads_once_and_verifies_real_parser_marker(tmp_path):
    from docx import Document
    calls = []
    extracted = ""
    class ParsingClient:
        async def upload(self, path):
            nonlocal extracted
            calls.append(path.name)
            extracted = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
            return {"id": "synthetic-file"}
        async def wait_for_file(self, file_id, *, timeout):
            assert file_id == "synthetic-file"
            return {"status": "ready", "text": extracted}
    options = args(tmp_path)
    options.file_type = "docx"
    result = asyncio.run(profile.workflow(ParsingClient(), "files", options, 1))
    assert calls == ["synthetic-profile.docx"]
    assert result["status"] == "passed"
    assert result["marker_verified"] is True
    assert extracted.startswith("LOADTEST-SYNTHETIC-")
    assert "合成行 099" in extracted
    assert "LOADTEST-SYNTHETIC-" not in json.dumps(result)
